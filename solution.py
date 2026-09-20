from __future__ import annotations

import argparse
import json
import os
import re
import time
from pathlib import Path
from typing import Any

from dotenv import load_dotenv

load_dotenv()


COMPARE_FIELDS = (
    "shipper",
    "consignee",
    "notify_party",
    "port_of_loading",
    "port_of_discharge",
    "container_count",
    "gross_weight_kg",
)

FIELD_PATTERNS = {
    "shipper": (r"shipper\s*\(principal or seller\)", r"shipper", r"发货人"),
    "consignee": (r"consignee\s*\(non-negotiable\)", r"consignee", r"to the order of", r"收货人"),
    "notify_party": (r"notify party\s*/\s*intermediate consignee", r"notify party", r"notify", r"通知人"),
    "port_of_loading": (r"port of loading\s*\(pol\)", r"port of loading", r"load port", r"pol", r"装货港"),
    "port_of_discharge": (r"port of discharge\s*\(pod\)", r"port of discharge", r"discharge port", r"pod", r"卸货港"),
    "container_count": (r"total containers", r"no\. of containers or packages", r"no\. of containers", r"container count", r"箱数"),
    "gross_weight_kg": (r"gross weight[^:|]*\(kgs\)", r"gross weight[^:|]*\(kg\)", r"gross wt[^:|]*\(kgs\)", r"gross weight", r"毛重[^:|]*kgs"),
}

BLANK_TOKENS = ("???", "_______", "TBA", "TBC", "N/A", "NA", "")
SPAM_MARKERS = ("GIFT CARD", "CLAIM NOW", "PARCEL IS ON HOLD", "STORAGE IS FULL", "90% OFF", "BITCOIN")
INVOICE_MARKERS = ("BILLING", "MISSING GR", "CANCEL INVOICE", "LOCAL CHARGES", "D & D CHARGES")
GENERAL_MARKERS = ("UPDATE SUMMARY", "BERTHING REPORT", "REMINDER", "RPA", "OUTSTANDING BL")
WRONG_DOC_MARKERS = ("COMMERCIAL INVOICE", "PACKING LIST", "CERTIFICATE OF ORIGIN")


# ---------------------------------------------------------------------------
# Optional AI extraction layer
# ---------------------------------------------------------------------------
# Falls back to pure regex extraction whenever no API key is configured, the
# dependency isn't installed, or the API call fails for any reason. This
# means the pipeline behaves identically to the regex-only baseline unless
# GEMINI_API_KEY (or GOOGLE_API_KEY) is actually set and working.

def get_gemini_client() -> Any | None:
    """Return a configured Gemini client when the API key is available."""
    api_key = os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")
    if not api_key:
        return None
    try:
        from google import genai
        return genai.Client(api_key=api_key)
    except Exception as exc:
        print(f"[extract_fields_ai] Gemini client unavailable: {exc}")
        return None


def extract_fields_ai(text: str) -> dict[str, str]:
    """AI-assisted extraction with evidence verification, falling back to regex."""
    client = get_gemini_client()
    if client is None:
        return extract_fields(text)

    field_list = ", ".join(COMPARE_FIELDS)
    prompt = (
        f"Extract these fields from the shipping document text: {field_list}. "
        f"For each field found, return the exact value AND the exact line of source "
        f"text it came from (verbatim, copy-pasted). Respond as JSON only: "
        f'{{"field_name": {{"value": "...", "evidence": "..."}}}}. '
        f"Omit fields you cannot find.\n\nDocument:\n{text[:6000]}"
    )
    try:
        response = client.models.generate_content(
            model="gemini-3.6-flash",
            contents=prompt,
            config={"response_mime_type": "application/json"},
        )
        parsed = json.loads(response.text)
    except Exception as exc:
        print(f"[extract_fields_ai] AI extraction failed, falling back to regex: {exc}")
        return extract_fields(text)

    fields: dict[str, str] = {}
    for field, data in parsed.items():
        if field not in COMPARE_FIELDS or not isinstance(data, dict):
            continue
        value = data.get("value", "")
        evidence = data.get("evidence", "")
        if evidence and evidence not in text:
            continue  # evidence doesn't actually exist in the source — reject it
        if value:
            fields[field] = value

    # Fill in anything the model missed using the regex path, rather than
    # leaving fields blank when AI extraction only partially succeeds.
    regex_fallback = extract_fields(text)
    for field, value in regex_fallback.items():
        fields.setdefault(field, value)
    return fields


def classify_email(subject: str, body: str, attachments: list[str]) -> str:
    del body
    text = subject.upper()
    if any(marker in text for marker in SPAM_MARKERS):
        return "SPAM"
    if any(marker in text for marker in INVOICE_MARKERS):
        return "INVOICE_QUERY"
    if re.search(r"\b(SI\s*-|REQUEST SI|CUST SI|SI NEEDED|LATEST SI)\b", text):
        return "SI_REQUEST"
    if any(marker in text for marker in GENERAL_MARKERS):
        return "GENERAL"
    if any(marker in text for marker in ("TO CONFIRM DOCS", "REQUEST BL DRAFT", "DRAFT BL")):
        return "BL_COMPARISON"
    return "BL_COMPARISON" if attachments else "GENERAL"


def read_document(path: Path) -> str:
    if not path.exists() or not path.is_file():
        return ""
    try:
        if path.suffix.lower() in {".txt", ".csv"}:
            return path.read_text(encoding="utf-8", errors="replace")
        if path.suffix.lower() == ".pdf":
            from pypdf import PdfReader
            return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
        if path.suffix.lower() == ".docx":
            from docx import Document
            document = Document(str(path))
            parts = [paragraph.text for paragraph in document.paragraphs]
            parts.extend(" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
            return "\n".join(parts)
        if path.suffix.lower() in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            rows = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).strip() for value in row if value is not None]
                    if values:
                        rows.append(" | ".join(values))
            return "\n".join(rows)
    except Exception:
        return ""
    return path.read_text(encoding="utf-8", errors="replace")


def extract_fields(text: str) -> dict[str, str]:
    fields: dict[str, str] = {}
    for field, patterns in FIELD_PATTERNS.items():
        label = "|".join(f"(?:{pattern})" for pattern in patterns)
        matcher = re.compile(rf"^\s*(?:{label})\s*[:|]\s*(.*?)\s*$", re.IGNORECASE)
        for line in text.splitlines():
            match = matcher.match(line.strip())
            if match:
                fields[field] = match.group(1).split(";")[0].strip()
                break
    return fields


def normalize_value(field: str, value: Any) -> Any:
    raw = str(value or "").strip()
    if field == "container_count":
        match = re.search(r"\d+", raw)
        return int(match.group()) if match else raw.upper()
    if field == "gross_weight_kg":
        match = re.search(r"[\d,]+(?:\.\d+)?", raw)
        if match:
            number = match.group().replace(",", "")
            return float(number) if "." in number else int(number)
        return raw.upper()
    value = re.sub(r"\s*\([A-Z0-9]+\)", "", raw.upper())
    value = re.sub(r"[,./_-]", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def has_blank_value(value: str | None) -> bool:
    normalized = str(value or "").strip().upper()
    if not normalized:
        return True
    return any(
        normalized == token or re.search(rf"(?<![A-Z0-9]){re.escape(token)}(?![A-Z0-9])", normalized)
        for token in BLANK_TOKENS[:-1]
    )


def audit_email(record: dict[str, Any], data_root: Path) -> dict[str, Any]:
    attachments = record.get("attachments") or []
    category = classify_email(record.get("subject", ""), record.get("body", ""), attachments)
    result = {"category": category, "status": "OK", "review_reason": None, "defect_fields": [], "has_defect": False}
    if category != "BL_COMPARISON":
        return result
    if len(attachments) < 2:
        result.update(status="NEEDS_REVIEW", review_reason="missing_attachment")
        return result

    si_text = read_document(data_root / attachments[0])
    bl_text = read_document(data_root / attachments[1])
    if len(si_text.strip()) < 10 or len(bl_text.strip()) < 10:
        result.update(status="NEEDS_REVIEW", review_reason="unreadable")
        return result
    # BUG FIX: previously only checked bl_text (attachments[1]) for the wrong
    # document type. Now checks both, since a misordered or misattached SI
    # could equally be the wrong document.
    if any(marker in si_text.upper() for marker in WRONG_DOC_MARKERS) or \
       any(marker in bl_text.upper() for marker in WRONG_DOC_MARKERS):
        result.update(status="NEEDS_REVIEW", review_reason="wrong_doc_type")
        return result

    si_fields = extract_fields_ai(si_text)
    bl_fields = extract_fields_ai(bl_text)
    if any(has_blank_value(si_fields.get(field)) or has_blank_value(bl_fields.get(field)) for field in COMPARE_FIELDS):
        result.update(status="NEEDS_REVIEW", review_reason="missing_value")
        return result

    defects = [
        field
        for field in COMPARE_FIELDS
        if normalize_value(field, si_fields.get(field)) != normalize_value(field, bl_fields.get(field))
    ]
    result.update(status="MISMATCH" if defects else "OK", defect_fields=sorted(defects), has_defect=bool(defects))
    return result


def build_submission(inbox_dir: Path, data_root: Path) -> dict[str, dict[str, Any]]:
    submission = {}
    for email_path in sorted(inbox_dir.glob("email_*.json")):
        record = json.loads(email_path.read_text(encoding="utf-8"))
        submission[record.get("email_id", email_path.stem)] = audit_email(record, data_root)
    return submission


def run_benchmark(root: Path) -> None:
    """Run reproducible parser stress tests and report measured throughput."""
    stress_cases = [
        ("port code normalization", normalize_value("port_of_loading", "Singapore (SGSIN)") == "SINGAPORE"),
        ("punctuation normalization", normalize_value("consignee", "East-Bright, FZ LLC") == "EAST BRIGHT FZ LLC"),
        ("container quantity parsing", normalize_value("container_count", "6 x 40'HC") == 6),
        ("weight comma parsing", normalize_value("gross_weight_kg", "131,058 KG") == 131058),
        ("blank marker detection", has_blank_value("???")),
        ("embedded NA is not blank", not has_blank_value("NANTONG")),
        ("spam routing", classify_email("BITCOIN CLAIM NOW", "", []) == "SPAM"),
        ("invoice routing", classify_email("LOCAL CHARGES QUERY", "", []) == "INVOICE_QUERY"),
    ]
    failed = [name for name, passed in stress_cases if not passed]
    started = time.perf_counter()
    submission = build_submission(root / "inbox", root)
    elapsed = time.perf_counter() - started
    throughput = len(submission) / elapsed if elapsed else float("inf")
    ai_active = get_gemini_client() is not None
    print(json.dumps({
        "records": len(submission),
        "elapsed_seconds": round(elapsed, 4),
        "records_per_second": round(throughput, 2),
        "adversarial_tests": len(stress_cases),
        "adversarial_passed": len(stress_cases) - len(failed),
        "adversarial_failed": failed,
        "ai_extraction_active": ai_active,
        "ground_truth_metrics": "unavailable: no labeled reference file is bundled",
    }, indent=2))
    if failed:
        raise AssertionError("benchmark failures: " + ", ".join(failed))


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate the shipping-document verification submission.")
    parser.add_argument("--root", type=Path, default=Path(__file__).parent)
    parser.add_argument("--output", type=Path, default=Path("submission.json"))
    parser.add_argument("--self-check", action="store_true")
    parser.add_argument("--benchmark", action="store_true")
    args = parser.parse_args()

    root = args.root.resolve()
    if args.benchmark:
        run_benchmark(root)
        return
    submission = build_submission(root / "inbox", root)
    args.output.write_text(json.dumps(submission, indent=2) + "\n", encoding="utf-8")
    counts: dict[str, int] = {}
    for result in submission.values():
        key = result["status"] if result["status"] == "NEEDS_REVIEW" else result["category"]
        counts[key] = counts.get(key, 0) + 1
    print(json.dumps({"emails": len(submission), "counts": counts, "output": str(args.output)}, indent=2))
    if args.self_check:
        assert len(submission) == 520, f"expected 520 emails, found {len(submission)}"
        assert submission["email_501"]["review_reason"] == "wrong_doc_type"
        assert submission["email_507"]["review_reason"] == "missing_attachment"
        assert submission["email_516"]["review_reason"] == "missing_value"
        print("self-check: passed")


if __name__ == "__main__":
    main()
